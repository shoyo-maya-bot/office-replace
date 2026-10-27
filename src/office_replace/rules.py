"""置換のルール（置換前 → 置換後）. 置換前には * と ? のワイルドカードが使える.

- `*` は 0 文字以上、`?` は 1 文字（Excel の検索と同じ）。`~*` `~?` `~~` で記号そのもの
  - `*` は Excel と同じく **できるだけ長く** 一致する（`見積*` は「見積書A」の全体に一致する）
  - ただし `*` と `?` は改行・タブをまたがず（Excel のセルの中の改行・タブの文字にも一致しない）、
    同じ置換前の 2 か所目の始まり（最初の `*` より前の部分）もまたがない.
    「株式会社*支店」は「株式会社A支店と株式会社B支店」の 2 か所に別々に一致する（間の「と」は残る）
  - 一致が長すぎる（LONG_MATCH 文字より長い）ときは、意図しない範囲かもしれないので警告を残す（long_matches）
- `--regex` のときは置換前を正規表現として扱い、置換後で \\1 などのグループを使える
- 対象（targets）: text（セル・図形・本文など）/ sheet（シート名）/ file（ファイル名）/ formula（Excel の数式の中の文字列）。既定は text
  - `all`（すべて）は text・sheet・file。formula は明示したときだけ（数式を書き換えるため）
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

TARGETS = ("text", "sheet", "file", "formula")
ALL_TARGETS = ("text", "sheet", "file")  # 「すべて」で選ばれるもの（formula は明示したときだけ）
_TARGET_ALIASES = {
    "文字列": "text",
    "テキスト": "text",
    "シート名": "sheet",
    "シート": "sheet",
    "ファイル名": "file",
    "ファイル": "file",
    "数式": "formula",
    "数式の文字列": "formula",
    "すべて": "all",
    "all": "all",
}


class _StringLoader(yaml.SafeLoader):
    """YAML の値を書いたとおりの文字列で読むローダー.

    safe_load は `0123` を 83（8 進数）、`no` を False、`1.0` を 1.0、`2026-01-01` を日付にするため、
    置換前・置換後が書いた文字と変わる. ここでは null（空・`~`・`null`）以外の暗黙の型の判定を外し、すべて文字列で読む.
    """


_StringLoader.yaml_implicit_resolvers = {
    ch: [(tag, rx) for tag, rx in resolvers if tag in ("tag:yaml.org,2002:null", "tag:yaml.org,2002:merge")]
    for ch, resolvers in yaml.SafeLoader.yaml_implicit_resolvers.items()
}


def _load_yaml_strings(text: str):
    """_StringLoader（SafeLoader の派生. 任意のオブジェクトは作らない）で YAML を読む."""
    loader = _StringLoader(text)
    try:
        return loader.get_single_data()
    finally:
        loader.dispose()


LONG_MATCH = 60  # * の一致がこの文字数を超えたら警告する


def _tokens(pattern: str) -> list[str]:
    """ワイルドカードの字句（"*" / "?" / 普通の 1 文字。~ の後の記号は普通の文字）."""
    out, i = [], 0
    while i < len(pattern):
        ch = pattern[i]
        if ch == "~" and i + 1 < len(pattern) and pattern[i + 1] in "*?~":
            out.append("~" + pattern[i + 1])
            i += 2
            continue
        out.append(ch)
        i += 1
    return out


# ワイルドカードが一致する 1 文字: 改行（\n・\r）とタブ以外. Word・PowerPoint の改行・タブは別のランとして切るが、
# Excel のセルの文字列には改行・タブがそのまま入るため、ここでも止める（README「タブ・改行はまたがない」）
_ANY = r"[^\n\r\t]"


def _piece(tok: str) -> str:
    if tok == "?":
        return _ANY
    return re.escape(tok[1:] if tok.startswith("~") else tok)


def wildcard_to_regex(pattern: str) -> str:
    """* と ? を正規表現にする. * も ? も改行・タブをまたがない（_ANY）.

    * は最長一致（Excel と同じ）だが、置換前の頭（最初の * より前）が次に現れる位置はまたがない.
    これで 1 段落・1 セルに 2 か所あるとき、1 か所目から 2 か所目の終わりまでを 1 つにまとめて消すことがない.
    """
    toks = _tokens(pattern)
    head_toks = toks[: toks.index("*")] if "*" in toks else []
    # 頭が ? だけ（どの文字にも一致する）なら、止める目印にならないので普通の最長一致
    head = "".join(_piece(t) for t in head_toks) if any(t != "?" for t in head_toks) else ""
    star = f"(?:(?!{head}){_ANY})*" if head else f"{_ANY}*"
    return "".join(star if t == "*" else _piece(t) for t in toks)


@dataclass
class Rule:
    before: str
    after: str
    targets: tuple[str, ...] = ("text",)
    regex: bool = False
    ignore_case: bool = False
    whole: bool = False  # セル・図形の文字列全体が一致するときだけ
    _re: re.Pattern[str] | None = field(default=None, repr=False)
    long_matches: list[int] = field(default_factory=list, repr=False)  # LONG_MATCH を超えた一致の長さ（警告用）

    def compiled(self) -> re.Pattern[str]:
        if self._re is None:
            body = self.before if self.regex else wildcard_to_regex(self.before)
            if self.whole:
                body = rf"\A(?:{body})\Z"
            flags = re.I if self.ignore_case else 0  # DOTALL は付けない（段落・セル内の改行で止める）
            self._re = re.compile(body, flags)
        return self._re

    def note(self, m: re.Match[str]) -> None:
        """ワイルドカード * の一致が長すぎれば記録する（置換はする. 警告としてレポートに出す）."""
        if not self.regex and "*" in _tokens(self.before) and m.end() - m.start() > LONG_MATCH:
            self.long_matches.append(m.end() - m.start())

    def apply(self, text: str) -> tuple[str, int]:
        rx = self.compiled()

        def rep(m: re.Match[str]) -> str:
            self.note(m)
            return m.expand(self.after) if self.regex else self.after

        return rx.subn(rep, text)


def validate(rules: list[Rule]) -> list[str]:
    problems = []
    for i, r in enumerate(rules, start=1):
        if not r.before:
            problems.append(f"ルール {i}: 置換前が空です")
            continue
        if not r.regex and r.before.strip("*") == "":
            problems.append(f"ルール {i}: 置換前が * だけでは、すべての文字列が置換されます")
        try:
            rx = r.compiled()
        except re.error as e:
            problems.append(f"ルール {i}: 正規表現の誤り: {e}")
            continue
        if rx.search("") and not r.whole:
            problems.append(f"ルール {i}: 空文字に一致するパターンです（{r.before}）")
        bad = [t for t in r.targets if t not in TARGETS]
        if bad:
            problems.append(f"ルール {i}: 対象 {bad} は text / sheet / file / formula のどれかです")
    return problems


def _targets(value: str) -> tuple[str, ...]:
    if not value:
        return ("text",)
    parts = [p.strip() for p in re.split(r"[,、/\s]+", value) if p.strip()]
    out: list[str] = []
    for p in parts:
        t = _TARGET_ALIASES.get(p, _TARGET_ALIASES.get(p.lower(), p.lower()))
        out += list(ALL_TARGETS) if t == "all" else [t]
    return tuple(dict.fromkeys(out))


def _row_rule(row: dict[str, str], *, regex: bool, ignore_case: bool) -> Rule | None:
    def get(*names: str) -> str:
        for n in names:
            if n in row and row[n] is not None:
                v = row[n]
                if isinstance(v, list):  # YAML の 対象: [text, sheet]
                    return ",".join(str(x) for x in v)
                return str(v)
        return ""

    before = get("置換前", "before", "from", "検索")
    if not before.strip():
        return None
    return Rule(
        before=before,
        after=get("置換後", "after", "to", "置換"),
        targets=_targets(get("対象", "targets", "target")),
        regex=regex,
        ignore_case=ignore_case,
        whole=get("完全一致", "whole").strip().lower() in ("○", "yes", "true", "1", "y"),
    )


_YAML_SHAPE = "rules: の下に「- 置換前: … / 置換後: …」の形のリストで書いてください"


def _yaml_rows(name: str, data: object) -> list[dict[str, str]]:
    """YAML の中身をルールの行（dict）のリストにする. 形が違えば ValueError（置換表の誤り＝終了コード 2）."""
    rows = data.get("rules", data) if isinstance(data, dict) else data
    if not rows:  # 空のファイル・rules: だけ → 呼び出し側の「ルールがありません」
        return []
    if isinstance(rows, dict):
        # トップ（または rules:）が 1 件のマッピング. 1 件だけでもリストにする（- を付ける）
        raise ValueError(f"{name}: ルールがリストになっていません（{_YAML_SHAPE}）")
    if not isinstance(rows, list):
        raise ValueError(f"{name}: ルールの形が違います（{_YAML_SHAPE}）")
    for i, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"{name}: {i} 件目のルールが「置換前: …」の組になっていません（{_YAML_SHAPE}）")
    return rows


def load(path: Path, *, regex: bool = False, ignore_case: bool = False) -> list[Rule]:
    """ルールの一覧（.csv / .tsv / .xlsx / .yaml）. 列: 置換前・置換後・対象（任意）・完全一致（任意）."""
    suffix = path.suffix.lower()
    rows: list[dict[str, str]]
    if suffix in (".yaml", ".yml"):
        # 値は書いたとおりの文字列で読む（0123 を 83、no を False にしない）
        data = _load_yaml_strings(path.read_text(encoding="utf-8-sig")) or {}
        rows = _yaml_rows(path.name, data)
    elif suffix in (".csv", ".tsv", ".txt"):
        text = None
        for enc in ("utf-8-sig", "cp932"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError(f"{path.name}: 文字コードを判別できません（UTF-8 か Shift_JIS）")
        # 改行を含むセル（"A<改行>B"）を壊さないよう、行に分けずにそのまま渡す
        rows = list(csv.DictReader(io.StringIO(text, newline=""), delimiter="\t" if suffix != ".csv" else ","))
    elif suffix in (".xlsx", ".xlsm"):
        from devtools_common.excel import render_cell
        from openpyxl import load_workbook

        # 数値のセルは表示どおりの文字列で読む（表示形式 000 の 1 は「001」、1.0 は「1」）
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows()
            header = [render_cell(c)[0].strip() for c in next(it, ())]
            rows = [{h: render_cell(c)[0] for h, c in zip(header, r, strict=False)} for r in it]
        finally:
            wb.close()
    else:
        raise ValueError(f"ルールの形式に対応していません（.csv / .tsv / .xlsx / .yaml）: {path.name}")
    rules = [r for r in (_row_rule(row, regex=regex, ignore_case=ignore_case) for row in rows) if r]
    if not rules:
        raise ValueError(f"{path.name}: ルールがありません（列名は 置換前・置換後）")
    return rules
