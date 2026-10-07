"""置換の結果のレポート（Markdown / Excel）.

レポートには置換前の文字列（伏せたい固有名）がそのまま入る。共有するときは mask（--mask-report）で伏せる。
"""

from __future__ import annotations

import re
from collections.abc import Callable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import FileResult
from .rules import Rule

STATUS = {"changed": "置換あり", "unchanged": "置換なし", "skipped": "対象外", "error": "エラー"}
KIND = {"text": "文字列", "sheet": "シート名", "file": "ファイル名", "formula": "数式の文字列"}
MASK = "（伏せ字）"


def _plain(s: str) -> str:
    return s


def masker(rules: list[Rule]) -> Callable[[str], str]:
    """レポートの文字列から置換前の文字列を消す関数. すべてのルールを（対象に関係なく）当てて置換後の文字にする."""

    def mask(text: str) -> str:
        text = str(text)
        for r in rules:
            try:
                text, _ = r.apply(text)
            except (re.error, IndexError):  # 置換後のグループ（\\1 など）の誤りは、文字ごと伏せる
                return MASK
        return text

    return mask


def summary(results: list[FileResult]) -> dict[str, int]:
    out = {k: 0 for k in STATUS}
    for r in results:
        out[r.status] += 1
    out["置換件数"] = sum(c.count for r in results for c in r.changes)
    out["残り"] = sum(x.count for r in results for x in r.leftovers)
    return out


def _cell(s: str) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ")


def markdown(results: list[FileResult], *, dry_run: bool, src: Path, out: Path, mask: Callable[[str], str] | None = None) -> str:
    """mask を渡すと、置換前の列は「（伏せ字）」に、ファイル名・場所・メッセージは置換後の文字にして出す."""
    s = summary(results)
    m = mask or _plain
    before = (lambda _v: MASK) if mask else _plain
    lines = [
        "# 置換の結果" + ("（試行: ファイルは書き換えていません）" if dry_run else ""),
        "",
        f"- 入力: {m(str(src))}",
        f"- 出力: {m(str(out)) if not dry_run else '（試行のため書き出しなし）'}",
        f"- ファイル: 置換あり {s['changed']} / 置換なし {s['unchanged']} / 対象外 {s['skipped']} / エラー {s['error']}",
        f"- 置換: {s['置換件数']} 件",
        f"- 置換しきれずに残った置換前の文字列: {s['残り']} 件" + ("（下の「残っている文字列」を確認してください）" if s["残り"] else ""),
        "",
    ]
    errs = [(r, msg) for r in results for msg in r.messages if r.status in ("error", "changed")]
    if errs:
        lines += ["## エラー・注意", "", "| ファイル | 内容 |", "|---|---|"]
        lines += [f"| {_cell(m(r.source.name))} | {_cell(m(msg))} |" for r, msg in errs]
        lines.append("")
    rows = [(r, c) for r in results for c in r.changes]
    if rows:
        lines += ["## 置換した箇所", "", "| ファイル | 種類 | 場所 | 置換前 | 置換後 | 件数 |", "|---|---|---|---|---|---|"]
        lines += [
            f"| {_cell(m(r.source.name))} | {KIND.get(c.kind, c.kind)} | {_cell(m(c.location))} | {_cell(before(c.before))} "
            f"| {_cell(m(c.after))} | {c.count} |"
            for r, c in rows
        ]
        lines.append("")
    left = [(r, x) for r in results for x in r.leftovers]
    if left:
        lines += [
            "## 残っている文字列（注意）",
            "",
            "置換後のファイルの中に、置換前の文字列がまだあります（このツールが置換しない場所）。「理由」の対処をしてください。",
            "固有名を伏せる目的のときは、Office で開いて確認してください。残りがあると終了コードは 1 です（`--allow-leftover` で 0）。",
            "",
            "| ファイル | 種類 | 部品 | 場所 | 置換前 | 件数 | 理由 |",
            "|---|---|---|---|---|---|---|",
        ]
        lines += [
            f"| {_cell(m(r.source.name))} | {x.rule_id} | {_cell(m(x.part))} | {_cell(x.where)} | {_cell(before(x.before))} | {x.count} "
            f"| {_cell(x.reason)} |"
            for r, x in left
        ]
        lines.append("")
    skipped = [r for r in results if r.status == "skipped" and r.messages and "古い形式" in r.messages[0]]
    if skipped:
        lines += ["## 対象外（古い形式）", ""] + [f"- {m(r.source.name)}: {r.messages[0]}" for r in skipped] + [""]
    return "\n".join(lines)


def excel(results: list[FileResult], path: Path, mask: Callable[[str], str] | None = None) -> None:
    m = mask or _plain
    before = (lambda _v: MASK) if mask else _plain
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:
        ws.append(header)
        for i, w in enumerate(widths, start=1):
            ws.cell(1, i).font = Font(bold=True)
            ws.cell(1, i).fill = head_fill
            ws.column_dimensions[get_column_letter(i)].width = w
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "ファイル"
    sheet(
        ws,
        ["ファイル", "出力先", "結果", "置換件数", "メッセージ"],
        [
            [m(str(r.source)), m(str(r.dest or "")), STATUS[r.status], sum(c.count for c in r.changes), m(" / ".join(r.messages))]
            for r in results
        ],
        [50, 50, 10, 10, 60],
    )
    sheet(
        wb.create_sheet("置換箇所"),
        ["ファイル", "種類", "場所", "置換前", "置換後", "件数"],
        [
            [m(r.source.name), KIND.get(c.kind, c.kind), m(c.location), before(c.before), m(c.after), c.count]
            for r in results
            for c in r.changes
        ],
        [30, 10, 30, 50, 50, 8],
    )
    sheet(
        wb.create_sheet("残り"),
        ["ファイル", "種類", "部品", "場所", "置換前", "件数", "理由"],
        [[m(r.source.name), x.rule_id, m(x.part), x.where, before(x.before), x.count, x.reason] for r in results for x in r.leftovers],
        [30, 18, 40, 30, 30, 8, 60],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
