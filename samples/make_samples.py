"""サンプルの入力（架空の会社「サンプル商事」向けの資料）と置換表を生成する.

    python samples/make_samples.py            # samples/input/ と samples/rules.* を作る

- input/見積書_サンプル商事.xlsx   セル・図形（テキストボックス）・シート名・シートを参照する数式
- input/提案書_サンプル商事.docx   本文（書式の違うランに分かれた文字列を含む）・表・ヘッダー
- input/sub/説明資料_サンプル商事.pptx  スライドの図形・表・ノート
- input/旧形式.xls                 古い形式（対象外として報告される）
- rules.csv / rules.xlsx           置換表（置換前・置換後・対象・完全一致）
"""

from __future__ import annotations

import csv
import shutil
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent

RULES = [
    {"置換前": "サンプル商事", "置換後": "〇〇商事", "対象": "すべて", "完全一致": ""},
    {"置換前": "株式会社テスト*支店", "置換後": "（支店名）", "対象": "文字列", "完全一致": ""},
    {"置換前": "担当: ?田", "置換後": "担当: （担当者）", "対象": "文字列", "完全一致": ""},
    {"置換前": "2026年度", "置換後": "2027年度", "対象": "文字列", "完全一致": ""},
]

DRAWING = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<xdr:twoCellAnchor><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>8</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="注記"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>
<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2400000" cy="800000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>
<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>サンプル</a:t></a:r><a:r><a:rPr b="1"/><a:t>商事</a:t></a:r><a:r><a:t> 様専用の価格です</a:t></a:r></a:p></xdr:txBody>
</xdr:sp><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>"""


def add_textbox_to_xlsx(path: Path, sheet_xml: str = "xl/worksheets/sheet1.xml") -> None:
    """openpyxl では図形を作れないので、zip に図形（テキストボックス）の XML を直接足す."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        names = src.namelist()
        for n in names:
            data = src.read(n)
            if n == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
                )
            elif n == sheet_xml:
                data = data.replace(
                    b"</worksheet>",
                    b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdDrw1"/></worksheet>',
                )
            out.writestr(n, data)
        rels = "xl/worksheets/_rels/sheet1.xml.rels"
        rel = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdDrw1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/>'
            "</Relationships>"
        )
        assert rels not in names
        out.writestr(rels, rel)
        out.writestr("xl/drawings/drawing1.xml", DRAWING)
    tmp.replace(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "サンプル商事向け"
    ws["A1"] = "御見積書"
    ws["A2"] = "サンプル商事 御中"
    ws["A3"] = "株式会社テスト東京支店"
    ws["A4"] = "担当: 山田"
    ws["A6"], ws["B6"] = "品目", "金額"
    ws["A7"], ws["B7"] = "設計（2026年度）", 1200000
    ws["A8"], ws["B8"] = "製造（2026年度）", 3400000
    ws["A9"], ws["B9"] = "合計", "=SUM(B7:B8)"
    summary = wb.create_sheet("集計")
    summary["A1"] = "見積合計"
    summary["B1"] = "='サンプル商事向け'!B9"
    summary["A2"] = "サンプル商事"
    wb.save(path)
    add_textbox_to_xlsx(path)


def make_docx(path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "サンプル商事様 提案書（社外秘）"
    doc.add_heading("サンプル商事様向け 業務改善のご提案", level=1)
    p = doc.add_paragraph()
    p.add_run("本書は")
    p.add_run("サンプル").bold = True
    p.add_run("商事様の 2026年度 の業務改善についての提案です。")
    doc.add_paragraph("担当: 田中（株式会社テスト大阪支店）")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "項目", "内容"
    t.cell(1, 0).text, t.cell(1, 1).text = "対象", "サンプル商事の受注業務"
    doc.save(path)


def make_pptx(path: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "サンプル商事様 ご説明資料"
    s1.placeholders[1].text = "2026年度 システム更改"
    s1.notes_slide.notes_text_frame.text = "サンプル商事の部長様が出席予定"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "体制"
    tb = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    para = tb.text_frame.paragraphs[0]
    r1 = para.add_run()
    r1.text = "窓口: サンプル"
    r2 = para.add_run()
    r2.text = "商事 情報システム部"
    r2.font.size = Pt(24)
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(3.5), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "役割", "担当"
    table.cell(1, 0).text, table.cell(1, 1).text = "営業", "株式会社テスト名古屋支店"
    prs.save(path)


def generate(out: Path) -> list[Path]:
    inp = out / "input"
    if inp.exists():
        shutil.rmtree(inp)
    (inp / "sub").mkdir(parents=True)
    files = [inp / "見積書_サンプル商事.xlsx", inp / "提案書_サンプル商事.docx", inp / "sub/説明資料_サンプル商事.pptx", inp / "旧形式.xls"]
    make_xlsx(files[0])
    make_docx(files[1])
    make_pptx(files[2])
    files[3].write_bytes(b"\xd0\xcf\x11\xe0 legacy binary placeholder")
    with (out / "rules.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(RULES[0]))
        w.writeheader()
        w.writerows(RULES)
    wb = Workbook()
    ws = wb.active
    ws.title = "置換表"
    ws.append(list(RULES[0]))
    for r in RULES:
        ws.append(list(r.values()))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24
    wb.save(out / "rules.xlsx")
    return [*files, out / "rules.csv", out / "rules.xlsx"]


if __name__ == "__main__":
    for p in generate(HERE):
        print(p)
