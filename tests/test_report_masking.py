"""--mask-report の伏せ字: 完全一致（whole）のルールでも、ファイル名・場所・エラー文の一部に入った置換前を伏せる."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from test_office_replace import report_of, run

from office_replace.report import masker
from office_replace.rules import Rule


def test_masker_applies_whole_rule_as_partial_match():
    mask = masker([Rule("架空商会", "〇〇", whole=True)])
    assert mask("架空商会_見積.xlsx") == "〇〇_見積.xlsx"
    assert mask("シート「架空商会向け」の A1") == "シート「〇〇向け」の A1"


def test_masker_does_not_change_the_rule_used_for_replacing():
    rule = Rule("架空商会", "〇〇", whole=True)
    masker([rule])("架空商会_見積")
    assert rule.apply("架空商会_見積") == ("架空商会_見積", 0)  # 置換には完全一致のまま


def _rules(path: Path) -> Path:
    path.write_text("置換前,置換後,対象,完全一致\n架空商会,〇〇,すべて,○\n", encoding="utf-8")
    return path


def test_mask_report_hides_whole_rule_in_file_names_and_messages(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active["A1"] = "架空商会"  # 完全一致で置換される
    wb.active["A2"] = "架空商会 御中"  # 完全一致では置換されない（レポートには出ない）
    wb.save(src / "架空商会_見積.xlsx")  # ファイル名は完全一致しないので変わらない
    (src / "架空商会_壊れ.docx").write_text("not a zip", encoding="utf-8")  # エラー文にファイル名が入る
    out = tmp_path / "out"
    rc, msg = run(str(src), "--rules", str(_rules(tmp_path / "rules.csv")), "-o", str(out), "--mask-report")
    assert rc == 1, msg  # 壊れたファイルがある
    md = report_of(out).read_text(encoding="utf-8")
    assert "架空商会" not in md, md
    assert "〇〇_見積.xlsx" in md and "〇〇_壊れ.docx" in md
    rep = load_workbook(report_of(out).with_suffix(".xlsx"))
    assert all("架空商会" not in str(c.value) for ws in rep.worksheets for row in ws.iter_rows() for c in row)
