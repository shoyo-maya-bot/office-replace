"""office-replace のテスト. サンプル生成スクリプトで入力を作り、置換後のファイルを各ライブラリで開いて確かめる."""

from __future__ import annotations

import contextlib
import importlib.util
import io
import zipfile
from pathlib import Path

import pytest
from devtools_common.cli import run_main
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from office_replace.cli import main
from office_replace.rules import Rule, load, validate, wildcard_to_regex
from office_replace.sheetname import check_sheet_name
from office_replace.textruns import replace_runs

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("make_samples", ROOT / "samples/make_samples.py")
assert _spec and _spec.loader
make_samples = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(make_samples)


@pytest.fixture()
def samples(tmp_path: Path) -> Path:
    make_samples.generate(tmp_path / "s")
    return tmp_path / "s"


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(main, list(args))
    return rc, out.getvalue() + err.getvalue()


# --- ルール ----------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("pattern", "text", "hit"),
    [
        ("会員*管理", "会員情報管理", True),
        ("担当: ?田", "担当: 山田", True),
        ("担当: ?田", "担当: 田中", False),
        ("100~*", "100*", True),
        ("100~*", "1000", False),
        ("a.b", "axb", False),  # ワイルドカード以外の記号はそのまま
    ],
)
def test_wildcard(pattern, text, hit):
    import re

    assert bool(re.search(wildcard_to_regex(pattern), text)) is hit


def test_rule_regex_groups_and_whole():
    assert Rule(r"第(\d+)版", r"Ver.\1", regex=True).apply("第12版") == ("Ver.12", 1)
    assert Rule("商事", "X", whole=True).apply("サンプル商事") == ("サンプル商事", 0)
    assert Rule("商事", "X", whole=True).apply("商事") == ("X", 1)
    assert Rule("abc", "X", ignore_case=True).apply("ABC") == ("X", 1)


def test_validate():
    assert validate([Rule("*", "x")])  # すべてに一致
    assert validate([Rule("(", "x", regex=True)])
    assert validate([Rule("a", "b", targets=("cell",))])
    assert validate([Rule("x*", "y")]) == []


def test_load_rules_csv_xlsx_cp932(samples: Path, tmp_path: Path):
    for f in ("rules.csv", "rules.xlsx"):
        rules = load(samples / f)
        assert [r.targets for r in rules][:2] == [("text", "sheet", "file"), ("text",)]
    p = tmp_path / "sjis.csv"
    p.write_text("置換前,置換後,対象,完全一致\n旧,新,シート名,○\n", encoding="cp932")
    r = load(p)[0]
    assert (r.before, r.after, r.targets, r.whole) == ("旧", "新", ("sheet",), True)
    (tmp_path / "empty.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError, match="置換前"):
        load(tmp_path / "empty.csv")


def test_replace_runs_across_formatting():
    assert replace_runs(["会員", "管理", "システム"], Rule("会員管理", "顧客管理")) == (["顧客管理", "", "システム"], 1)
    assert replace_runs(["ab", "cd", "ef"], Rule("b*e", "-")) == (["a-", "", "f"], 1)
    assert replace_runs(["xx", "yy"], Rule("x", "Z")) == (["ZZ", "yy"], 2)


@pytest.mark.parametrize(
    ("name", "problem"),
    [
        ("a" * 32, "31 文字"),
        ("集計[1]", "使えない文字"),
        ("a/b", "使えない文字"),
        ("History", "予約語"),
        ("集計", "同じ名前"),
        ("'x", "アポストロフィ"),
        ("", "空"),
    ],
)
def test_sheet_name_checks(name, problem):
    assert problem in check_sheet_name(name, ["集計"])


# --- 置換の実行 --------------------------------------------------------------------------------


def test_replace_folder(samples: Path, tmp_path: Path):
    out = tmp_path / "out"
    rc, msg = run(str(samples / "input"), "--rules", str(samples / "rules.xlsx"), "-r", "-o", str(out))
    assert rc == 0, msg
    # Excel: セル・図形・シート名・シートを参照する数式
    wb = load_workbook(out / "見積書_〇〇商事.xlsx")
    assert wb.sheetnames == ["〇〇商事向け", "集計"]
    ws = wb["〇〇商事向け"]
    assert [ws["A2"].value, ws["A3"].value, ws["A4"].value, ws["A7"].value] == [
        "〇〇商事 御中",
        "（支店名）",
        "担当: （担当者）",
        "設計（2027年度）",
    ]
    assert ws["B7"].value == 1200000  # 数値は変えない
    assert wb["集計"]["B1"].value == "='〇〇商事向け'!B9"
    with zipfile.ZipFile(out / "見積書_〇〇商事.xlsx") as z:
        drawing = z.read("xl/drawings/drawing1.xml").decode()
    assert "<a:t>〇〇商事</a:t>" in drawing and 'b="1"' in drawing  # 書式付きのランに入る
    # Word: ランをまたぐ文字列・書式・ヘッダー・表
    doc = Document(out / "提案書_〇〇商事.docx")
    para = doc.paragraphs[1]
    assert para.text == "本書は〇〇商事様の 2027年度 の業務改善についての提案です。"
    assert para.runs[1].text == "〇〇商事" and para.runs[1].bold
    assert doc.sections[0].header.paragraphs[0].text == "〇〇商事様 提案書（社外秘）"
    assert doc.tables[0].cell(1, 1).text == "〇〇商事の受注業務"
    # PowerPoint（サブフォルダ）: 図形・表・ノート
    prs = Presentation(out / "sub/説明資料_〇〇商事.pptx")
    s1, s2 = prs.slides
    assert s1.shapes.title.text == "〇〇商事様 ご説明資料"
    assert s1.notes_slide.notes_text_frame.text == "〇〇商事の部長様が出席予定"
    texts = [sh.text_frame.text for sh in s2.shapes if sh.has_text_frame]
    assert "窓口: 〇〇商事 情報システム部" in texts
    table = next(sh for sh in s2.shapes if sh.has_table).table
    assert table.cell(1, 1).text == "（支店名）"
    # 元のファイルは変わらない・古い形式は出力しない
    assert load_workbook(samples / "input/見積書_サンプル商事.xlsx").sheetnames[0] == "サンプル商事向け"
    assert not (out / "旧形式.xls").exists()
    # レポート
    md = (out / "office-replace-report.md").read_text(encoding="utf-8")
    assert "シート名" in md and "古い形式" in md and "1 か所を新しいシート名に更新" in md
    assert load_workbook(out / "office-replace-report.xlsx").sheetnames == ["ファイル", "置換箇所", "残り"]


def test_unchanged_parts_are_copied_byte_for_byte(samples: Path, tmp_path: Path):
    out = tmp_path / "out"
    run(str(samples / "input"), "--from", "2026年度", "--to", "2027年度", "-o", str(out))
    with zipfile.ZipFile(samples / "input/提案書_サンプル商事.docx") as a, zipfile.ZipFile(out / "提案書_サンプル商事.docx") as b:
        assert a.namelist() == b.namelist()
        for n in a.namelist():
            if n != "word/document.xml":
                assert a.read(n) == b.read(n), n


def test_dry_run_writes_nothing(samples: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.chdir(tmp_path)
    rc, out = run(str(samples / "input"), "--rules", str(samples / "rules.csv"), "-r", "--dry-run")
    assert rc == 0
    assert "試行" in out and "| 見積書_サンプル商事.xlsx | シート名 |" in out
    assert not (tmp_path / "_dry_run").exists()


def test_invalid_sheet_name_is_skipped_and_reported(samples: Path, tmp_path: Path):
    out = tmp_path / "out"
    rc, msg = run(
        str(samples / "input/見積書_サンプル商事.xlsx"),
        "--from",
        "サンプル商事向け",
        "--to",
        "見積/2027",
        "--target",
        "sheet",
        "-o",
        str(out),
    )
    assert rc == 1
    assert "使えない文字" in (out / "office-replace-report.md").read_text(encoding="utf-8")


def test_sheet_name_too_long(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active.title = "短い"
    wb.save(src / "a.xlsx")
    rc, _ = run(str(src), "--from", "短い", "--to", "長" * 32, "--target", "sheet", "-o", str(tmp_path / "out"))
    assert rc == 1
    assert "31 文字" in (tmp_path / "out/office-replace-report.md").read_text(encoding="utf-8")


def test_file_name_collision(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    for n in ("資料A.docx", "資料B.docx"):
        Document().save(src / n)
    rc, _ = run(str(src), "--from", "資料?", "--to", "資料", "--target", "file", "-o", str(tmp_path / "out"))
    assert rc == 1
    assert len(list((tmp_path / "out").glob("*.docx"))) == 1


def test_output_inside_input_is_refused(samples: Path):
    rc, out = run(str(samples / "input"), "--from", "a", "--to", "b", "-o", str(samples / "input/out"))
    assert rc == 2 and "入力フォルダの中" in out


def test_broken_file_is_error(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    (src / "pw.xlsx").write_bytes(b"not a zip")
    rc, _ = run(str(src), "--from", "a", "--to", "b", "-o", str(tmp_path / "out"))
    assert rc == 1
    assert "開けません" in (tmp_path / "out/office-replace-report.md").read_text(encoding="utf-8")


def test_copy_all(samples: Path, tmp_path: Path):
    out = tmp_path / "out"
    run(str(samples / "input"), "--from", "存在しない語", "--to", "x", "-r", "--copy-all", "-o", str(out))
    assert (out / "提案書_サンプル商事.docx").exists() and (out / "sub/説明資料_サンプル商事.pptx").exists()


def test_save_failure_leaves_no_temp_file(tmp_path: Path, monkeypatch):
    from office_replace import ooxml

    src = tmp_path / "in.docx"
    d = Document()
    d.add_paragraph("x")
    d.save(src)
    pkg = ooxml.Package(src)
    dest = tmp_path / "out" / "in.docx"

    def boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(zipfile.ZipFile, "writestr", boom)
    with pytest.raises(OSError):
        pkg.save(dest)
    assert list(dest.parent.iterdir()) == []


@pytest.mark.parametrize("to", ["../../escaped", "sub/dir", "a:b", "  ", "末尾."])
def test_file_name_replace_cannot_escape_or_break(tmp_path: Path, to: str):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.save(src / "report_FY.xlsx")
    out = tmp_path / "out"
    rc, _ = run(str(src), "--from", "report_FY", "--to", to, "--target", "file", "-o", str(out))
    assert rc != 0
    written = [p for p in tmp_path.rglob("*.xlsx") if p.parent != src and not p.name.startswith("office-replace-report")]
    assert written == []  # 出力フォルダの外にも中にも書かない


# --- 固有名の残り（プロパティ・グラフ・ふりがな・コメント）と、シート名の参照 ------------------------------


def _patch_zip(path: Path, edits: dict[str, object]) -> None:
    """zip の部品を書き換える／足す. 値は bytes か str、または 元の bytes を受けて新しい bytes を返す関数."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        names = src.namelist()
        for n in names:
            data = src.read(n)
            e = edits.get(n)
            if callable(e):
                data = e(data)
            elif e is not None:
                data = e.encode() if isinstance(e, str) else e
            out.writestr(n, data)
        for n, e in edits.items():
            if n not in names:
                out.writestr(n, e(b"") if callable(e) else e)
    tmp.replace(path)


def _read(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def _all_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels", ".vml")))


SST_RPH = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="2" uniqueCount="2">'
    '<si><t>架空商会</t><rPh sb="0" eb="2"><t>カクウ</t></rPh><rPh sb="2" eb="4"><t>ショウカイ</t></rPh>'
    '<phoneticPr fontId="1"/></si>'
    '<si><t>無関係</t><rPh sb="0" eb="3"><t>ムカンケイ</t></rPh></si></sst>'
)

CHART = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart><c:plotArea><c:barChart><c:ser>'
    '<c:tx><c:strRef><c:f>Sheet1!$B$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>架空商会 売上</c:v></c:pt>'
    "</c:strCache></c:strRef></c:tx>"
    '<c:cat><c:strRef><c:f>Sheet1!$A$2</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>架空商会</c:v></c:pt>'
    "</c:strCache></c:strRef></c:cat>"
    '<c:val><c:numRef><c:f>Sheet1!$B$2</c:f><c:numCache><c:ptCount val="1"/><c:pt idx="0"><c:v>10</c:v></c:pt>'
    "</c:numCache></c:numRef></c:val></c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>"
)


def test_excel_properties_chart_phonetic_comments(tmp_path: Path):
    from openpyxl.comments import Comment

    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "見出し"
    ws["B2"] = "数値"
    ws["A1"].comment = Comment("架空商会の担当に確認", "架空 太郎")
    wb.properties.title = "架空商会 見積"
    wb.properties.subject = "架空商会向け"
    wb.properties.creator = "架空 太郎"
    wb.properties.keywords = "架空商会"
    wb.properties.description = "架空商会の資料"
    wb.properties.lastModifiedBy = "架空 太郎"
    path = src / "a.xlsx"
    wb.save(path)

    def app(data: bytes) -> bytes:
        return data.replace(b"</Properties>", "<Company>架空商会</Company><Manager>架空 部長</Manager></Properties>".encode())

    threaded = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
        '<threadedComment ref="B2" id="{00000000-0000-0000-0000-000000000001}" personId="{00000000-0000-0000-0000-000000000002}">'
        "<text>架空商会に送付済み</text></threadedComment></ThreadedComments>"
    )
    persons = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<personList xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
        '<person displayName="架空 太郎" id="{00000000-0000-0000-0000-000000000002}" userId="u1" providerId="None"/></personList>'
    )
    _patch_zip(
        path,
        {
            "docProps/app.xml": app,
            "xl/sharedStrings.xml": SST_RPH,
            "xl/charts/chart1.xml": CHART,
            "xl/threadedComments/threadedComment1.xml": threaded,
            "xl/persons/person.xml": persons,
        },
    )
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out))
    assert rc == 0, msg
    res = out / "a.xlsx"
    xml = _all_xml(res)
    assert "架空" not in xml  # どの部品にも残らない
    core = _read(res, "docProps/core.xml")
    assert "〇〇商会 見積" in core and "〇〇 太郎" in core
    assert "<Company>〇〇商会</Company>" in _read(res, "docProps/app.xml")
    chart = _read(res, "xl/charts/chart1.xml")
    assert "<c:v>〇〇商会</c:v>" in chart and "<c:v>〇〇商会 売上</c:v>" in chart and "<c:v>10</c:v>" in chart
    sst = _read(res, "xl/sharedStrings.xml")
    assert "カクウ" not in sst and "<t>〇〇商会</t>" in sst  # 置換した文字列のふりがなは外す
    assert "ムカンケイ" in sst  # 置換していない文字列のふりがなは残す
    assert "〇〇商会に送付済み" in _read(res, "xl/threadedComments/threadedComment1.xml")
    assert 'displayName="〇〇 太郎"' in _read(res, "xl/persons/person.xml")
    wb2 = load_workbook(res)
    assert wb2.active["A1"].comment.text.startswith("〇〇商会の担当")
    md = (out / "office-replace-report.md").read_text(encoding="utf-8")
    assert "プロパティ" in md and "コメント" in md and "replace.leftover" not in md


def test_word_properties_and_chart(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    d = Document()
    d.add_paragraph("本文")
    d.core_properties.author = "架空 花子"
    d.core_properties.title = "架空商会 報告"
    d.save(src / "a.docx")
    _patch_zip(src / "a.docx", {"word/charts/chart1.xml": CHART})
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out))
    assert rc == 0, msg
    assert "架空" not in _all_xml(out / "a.docx")
    assert Document(out / "a.docx").core_properties.author == "〇〇 花子"


def test_powerpoint_chart_cache_and_comments(tmp_path: Path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    src = tmp_path / "in"
    src.mkdir()
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "表紙"
    cd = CategoryChartData()
    cd.categories = ["架空商会", "その他"]
    cd.add_series("架空商会 売上", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), cd)
    prs.core_properties.author = "架空 次郎"
    path = src / "a.pptx"
    prs.save(path)
    legacy = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cm authorId="0" idx="1"><p:pos x="0" y="0"/><p:text>架空商会の数字を確認</p:text></p:cm></p:cmLst>'
    )
    authors = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cmAuthor id="0" name="架空 次郎" initials="架" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>'
    )
    modern = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p188:cmLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        '<p188:cm id="{00000000-0000-0000-0000-000000000003}" authorId="{00000000-0000-0000-0000-000000000004}">'
        "<p188:txBody><a:bodyPr/><a:p><a:r><a:t>架空</a:t></a:r><a:r><a:t>商会に共有</a:t></a:r></a:p></p188:txBody>"
        '<p188:replyLst><p188:reply id="{00000000-0000-0000-0000-000000000005}" authorId="{00000000-0000-0000-0000-000000000004}">'
        "<p188:txBody><a:bodyPr/><a:p><a:r><a:t>架空商会 了解</a:t></a:r></a:p></p188:txBody></p188:reply></p188:replyLst>"
        "</p188:cm></p188:cmLst>"
    )
    modern_authors = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p188:authorLst xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        '<p188:author id="{00000000-0000-0000-0000-000000000004}" name="架空 次郎" initials="KJ" userId="u1" providerId="None"/>'
        "</p188:authorLst>"
    )
    _patch_zip(
        path,
        {
            "ppt/comments/comment1.xml": legacy,
            "ppt/commentAuthors.xml": authors,
            "ppt/comments/modernComment_100_0.xml": modern,
            "ppt/authors.xml": modern_authors,
        },
    )
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out))
    assert rc == 0, msg
    res = out / "a.pptx"
    with zipfile.ZipFile(res) as z:
        chart_part = next(n for n in z.namelist() if n.startswith("ppt/charts/chart") and n.endswith(".xml"))
    chart = _read(res, chart_part)
    assert "<c:v>〇〇商会</c:v>" in chart and "<c:v>〇〇商会 売上</c:v>" in chart
    assert "〇〇商会の数字を確認" in _read(res, "ppt/comments/comment1.xml")
    assert 'name="〇〇 次郎"' in _read(res, "ppt/commentAuthors.xml")
    assert "〇〇商会に共有" in _read(res, "ppt/comments/modernComment_100_0.xml").replace("</a:t></a:r><a:r><a:t>", "")
    assert "〇〇商会 了解" in _read(res, "ppt/comments/modernComment_100_0.xml")
    assert 'name="〇〇 次郎"' in _read(res, "ppt/authors.xml")
    assert Presentation(res).core_properties.author == "〇〇 次郎"
    # グラフの元データ（埋め込みのブック）は置換しないので、残りとしてレポートに出る
    md = (out / "office-replace-report.md").read_text(encoding="utf-8")
    assert "replace.leftover" in md and "embeddings/" in md
    assert load_workbook(out / "office-replace-report.xlsx").sheetnames == ["ファイル", "置換箇所", "残り"]


def test_leftover_is_reported(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "架空商会"
    ws["A2"] = '="架空商会"&"様"'  # 数式の中の文字列は置換しない
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(out))
    assert rc == 0, msg  # 注意であり、エラーではない
    md = (out / "office-replace-report.md").read_text(encoding="utf-8")
    assert "replace.leftover" in md and "xl/worksheets/sheet1.xml" in md and "架空商会" in md
    # 置換後に置換前が含まれるルール（集計 → 集計_2027）は、残りと見なさない
    rc, _ = run(str(src), "--from", "架空", "--to", "架空X", "-o", str(tmp_path / "out2"))
    assert "replace.leftover" not in (tmp_path / "out2/office-replace-report.md").read_text(encoding="utf-8")
    # 試行でも出る
    rc, dry = run(str(src), "--from", "架空商会", "--to", "〇〇", "--dry-run")
    assert "replace.leftover" in dry


@pytest.mark.filterwarnings("ignore:.*extension is not supported:UserWarning")
def test_sheet_rename_updates_validation_and_conditional_formatting(tmp_path: Path):
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "入力"
    lst = wb.create_sheet("架空 一覧")
    lst["A1"], lst["A2"] = "甲", "乙"
    dv = DataValidation(type="list", formula1="'架空 一覧'!$A$1:$A$2")
    ws.add_data_validation(dv)
    dv.add("B1")
    dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="'架空 一覧'!$B$1")
    ws.add_data_validation(dv2)
    dv2.add("C1")
    ws.conditional_formatting.add("B1", FormulaRule(formula=["B1='架空 一覧'!$A$1"]))
    path = src / "a.xlsx"
    wb.save(path)
    x14 = (
        '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" count="1">'
        '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>\'架空 一覧\'!$A$1:$A$2</xm:f></x14:formula1>'
        "<xm:sqref>D1</xm:sqref></x14:dataValidation></x14:dataValidations></ext>"
        '<ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
        '<x14:cfRule type="expression" priority="2" id="{00000000-0000-0000-0000-000000000006}">'
        "<xm:f>D1='架空 一覧'!$A$2</xm:f></x14:cfRule><xm:sqref>D1</xm:sqref></x14:conditionalFormatting>"
        "</x14:conditionalFormattings></ext></extLst></worksheet>"
    )

    def sheet1(data: bytes) -> bytes:
        return data.replace(b"</worksheet>", x14.encode())

    _patch_zip(path, {"xl/worksheets/sheet1.xml": sheet1})
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空 一覧", "--to", "選択肢", "--target", "sheet", "-o", str(out))
    assert rc == 0, msg
    s1 = _read(out / "a.xlsx", "xl/worksheets/sheet1.xml")
    assert "架空" not in s1
    assert "<formula1>'選択肢'!$A$1:$A$2</formula1>" in s1
    assert "<formula2>'選択肢'!$B$1</formula2>" in s1
    assert "<formula>B1='選択肢'!$A$1</formula>" in s1
    assert "<xm:f>'選択肢'!$A$1:$A$2</xm:f>" in s1 and "<xm:f>D1='選択肢'!$A$2</xm:f>" in s1
    assert load_workbook(out / "a.xlsx").sheetnames == ["入力", "選択肢"]
    assert "架空" not in _read(out / "a.xlsx", "docProps/app.xml")  # 文書のプロパティのシート名の一覧も合わせる
