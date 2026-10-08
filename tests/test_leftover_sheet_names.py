"""残りの確認: シート名が原因で残る置換前の文字列には、--target sheet で置換できることを理由に書く."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from test_office_replace import report_of, run

from office_replace import leftover
from office_replace.processors.formula import sheet_ref_spans


def _book(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "架空商会"
    wb.active["A1"] = 100
    other = wb.create_sheet("集計")
    other["A1"] = "=架空商会!A1*2"  # シート名の参照だけ
    other["A2"] = "='架空商会'!A1"
    wb.save(path)


def _leftover_rows(md: str) -> list[str]:
    return [line for line in md.splitlines() if line.startswith("|") and "replace.leftover" in line]


def test_leftover_from_sheet_name_suggests_target_sheet(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _book(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(out))
    assert rc == 1, msg  # 残りの規則は変えない（残りがあれば 1）
    rows = _leftover_rows(report_of(out).read_text(encoding="utf-8"))
    sheet_row = next(r for r in rows if "xl/worksheets/sheet2.xml" in r)
    assert "--target sheet" in sheet_row and "--target formula" not in sheet_row
    book_row = next(r for r in rows if "xl/workbook.xml" in r)
    assert "sheet@name" in book_row and "--target sheet" in book_row
    # --target sheet を付ければ、シート名も参照も残らない
    out2 = tmp_path / "out2"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "--target", "text,sheet", "-o", str(out2))
    assert rc == 0, msg
    assert not _leftover_rows(report_of(out2).read_text(encoding="utf-8"))


def test_literal_in_formula_still_suggests_target_formula(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active["A1"] = '="架空商会"&"様"'
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(out))
    row = next(r for r in _leftover_rows(report_of(out).read_text(encoding="utf-8")) if "sheet1.xml" in r)
    assert "--target formula" in row


def test_reason_for_sheet_refs_and_sheet_name_attribute():
    assert "--target sheet" in leftover.reason("xl/worksheets/sheet1.xml", "f" + leftover.SHEET_REF)
    assert "--target sheet" in leftover.reason("xl/workbook.xml", "definedName" + leftover.SHEET_REF)
    assert "--target sheet" in leftover.reason("xl/workbook.xml", "sheet@name")
    assert "--target formula" in leftover.reason("xl/worksheets/sheet1.xml", "f")


def test_sheet_ref_spans_skip_strings_and_external_books():
    f = "架空!A1&\"架空!\"&[1]架空!A1&SUM(集計:架空!B1)&'架空 2'!C1"
    assert [f[a:b] for a, b in sheet_ref_spans(f)] == ["架空!", "集計:架空!", "'架空 2'!"]
