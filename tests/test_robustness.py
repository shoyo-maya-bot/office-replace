"""残りの判定（docProps の日付・Excel の v）・1 ファイルの失敗で止めない・zip 爆弾対策."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from test_office_replace import _patch_zip, report_of, run

from office_replace import engine, leftover
from office_replace.ooxml import Package, UnsafeZipError, ZipLimits, check_zip

CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
    ' xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"'
    ' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    "<dc:title>2025年度 計画</dc:title>"
    '<dcterms:created xsi:type="dcterms:W3CDTF">2025-04-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2025-05-01T00:00:00Z</dcterms:modified>'
    "<cp:lastPrinted>2025-05-02T00:00:00Z</cp:lastPrinted>"
    "</cp:coreProperties>"
)
CUSTOM = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"'
    ' xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="締切"><vt:filetime>2025-06-30T00:00:00Z</vt:filetime></property>'
    "</Properties>"
)


# --- 残りの判定 -------------------------------------------------------------------------------------


def test_year_replace_does_not_flag_docprops_dates(tmp_path: Path):
    """第 1 回の後退: 作成日・更新日（2025-…）が「残り」と判定され、年度の置換が常に終了コード 1 になっていた."""
    src = tmp_path / "in"
    src.mkdir()
    d = Document()
    d.add_paragraph("2025年度の計画")
    d.save(src / "a.docx")
    _patch_zip(src / "a.docx", {"docProps/core.xml": CORE})
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "2025", "--to", "2026", "-o", str(out))
    assert rc == 0, msg
    assert "replace.leftover" not in report_of(out).read_text(encoding="utf-8")
    with zipfile.ZipFile(out / "a.docx") as z:
        core = z.read("docProps/core.xml").decode()
    assert "2026年度 計画" in core  # タイトルは置換する
    assert "2025-04-01" in core  # 日付は置換もしない（文字列ではない値）


def test_docprops_dates_are_skipped_but_text_is_checked():
    strings = leftover._strings(CORE.encode(), None, "docProps/core.xml")
    assert ("title", "2025年度 計画") in strings
    assert not any(w in ("created", "modified", "lastPrinted") for w, _ in strings)
    assert not any(w == "filetime" for w, _ in leftover._strings(CUSTOM.encode(), None, "docProps/custom.xml"))
    # docProps 以外の同名の要素は調べる（日付だから除くのは docProps だけ）
    other = b'<r xmlns="urn:x"><created>2025</created></r>'
    assert ("created", "2025") in leftover._strings(other, None, "word/document.xml")


def test_excel_v_is_classified(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2025年度"  # 共有文字列（t="s" の v は番号なので調べない）
    ws["A2"] = 2025  # 数値のセル
    wb.save(src / "a.xlsx")
    # 数式の計算結果（t="str" の v）を足す
    _patch_zip(
        src / "a.xlsx",
        {
            "xl/worksheets/sheet1.xml": lambda b: b.replace(
                b"</sheetData>", b'<row r="3"><c r="A3" t="str"><f>"2025"&amp;"x"</f><v>2025x</v></c></row></sheetData>'
            )
        },
    )
    out = tmp_path / "out"
    rc, _ = run(str(src), "--from", "2025", "--to", "2026", "-o", str(out))
    assert rc == 1  # 数値・計算結果に残る（本当に残っている）
    md = report_of(out).read_text(encoding="utf-8")
    assert "v（数値など）" in md and "v（数式の結果）" in md
    assert "このツールは数値を置換しません" in md
    # v の説明は実際の挙動（計算結果は書き換えない）に合わせる
    assert "このツールは計算結果を書き換えません" in leftover.reason("xl/worksheets/sheet1.xml", "v（数式の結果）")


# --- 1 ファイルの失敗で止めない ---------------------------------------------------------------------


def test_one_file_failure_does_not_stop_the_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    src = tmp_path / "in"
    src.mkdir()
    for name in ("a.docx", "b.docx", "c.docx"):
        d = Document()
        d.add_paragraph("架空商会")
        d.save(src / name)
    real = engine.process_file

    def flaky(path: Path, *a, **k):
        if path.name == "b.docx":
            raise RuntimeError("想定外の XML")
        return real(path, *a, **k)

    monkeypatch.setattr(engine, "process_file", flaky)
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out))
    assert rc == 1, msg  # 失敗したファイルがあれば 1（エラーのあったファイルあり）
    assert (out / "a.docx").exists() and (out / "c.docx").exists() and not (out / "b.docx").exists()
    md = report_of(out).read_text(encoding="utf-8")  # レポートは必ず書く
    assert "RuntimeError: 想定外の XML" in md and "エラー 1" in md


# --- zip 爆弾対策 ----------------------------------------------------------------------------------


def _bomb(path: Path, size: int) -> None:
    Document().save(path)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/media/zeros.bin", b"\0" * size)


def test_zip_limits_reject_bomb(tmp_path: Path):
    p = tmp_path / "bomb.docx"
    _bomb(p, 8 * 1024 * 1024)  # 8MB のゼロ（圧縮率は約 1000 倍）
    with zipfile.ZipFile(p) as z:
        with pytest.raises(UnsafeZipError, match="圧縮率"):
            check_zip(z, ZipLimits())
        check_zip(z, ZipLimits(max_ratio=2000))  # 上限は変えられる
        with pytest.raises(UnsafeZipError, match="展開後の大きさ"):
            check_zip(z, ZipLimits(max_total=1024 * 1024, max_ratio=2000))
        with pytest.raises(UnsafeZipError, match="部品の数"):
            check_zip(z, ZipLimits(max_entries=3, max_ratio=2000))
    with pytest.raises(UnsafeZipError):
        Package(p)


def test_zip_bomb_is_error_and_run_continues(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _bomb(src / "bomb.docx", 8 * 1024 * 1024)
    d = Document()
    d.add_paragraph("架空商会")
    d.save(src / "ok.docx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out))
    assert rc == 1, msg
    assert (out / "ok.docx").exists() and not (out / "bomb.docx").exists()
    assert "安全のため処理しません" in report_of(out).read_text(encoding="utf-8")
    # CLI で上限を上げれば処理できる
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(tmp_path / "out2"), "--max-ratio", "5000")
    assert rc == 0, msg


def test_zip_limit_options_are_validated(tmp_path: Path):
    with pytest.raises(SystemExit) as e:
        run(str(tmp_path), "--from", "a", "--to", "b", "--dry-run", "--max-unzip-mb", "0")
    assert e.value.code == 2


def test_normal_office_files_are_within_default_limits(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    for i in range(1, 20001):  # 繰り返しの多い大きめのシート（普通の業務ファイル）
        ws.append([f"項目{i % 10}", i, "同じ文字列"])
    wb.save(tmp_path / "big.xlsx")
    with zipfile.ZipFile(tmp_path / "big.xlsx") as z:
        check_zip(z, ZipLimits())
