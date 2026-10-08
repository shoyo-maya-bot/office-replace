"""段落の中の改行（w:br・a:br）とタブ（w:tab）: 段落の文字列に入れ、置換（* を含む）はまたがない."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation
from test_office_replace import _read, run

from office_replace.rules import Rule
from office_replace.textruns import replace_runs


def test_replace_runs_does_not_cross_fixed_runs():
    # 1 番目が改行. 「A*B」は改行をまたいで一致しない
    assert replace_runs(["A1", "\n", "2B"], Rule("A*B", "Z"), {1}) == (["A1", "\n", "2B"], 0)
    assert replace_runs(["会員", "\t", "管理"], Rule("会員管理", "X"), {1}) == (["会員", "\t", "管理"], 0)
    # 改行の両側では別々に置換する（改行は残る）
    assert replace_runs(["A1B", "\n", "A2B"], Rule("A*B", "Z"), {1}) == (["Z", "\n", "Z"], 2)
    # 改行で区切られた範囲の中では、ラン（書式の切れ目）をまたいで置換する
    assert replace_runs(["x", "\n", "会", "員"], Rule("会員", "客"), {2 - 1}) == (["x", "\n", "客", ""], 1)


def _word(path: Path) -> None:
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("A1")
    r.add_break()  # w:br
    p.add_run("2B")
    p2 = d.add_paragraph()
    r2 = p2.add_run("会員")
    r2.add_tab()  # w:tab
    p2.add_run("管理")
    p3 = d.add_paragraph()
    r3 = p3.add_run("A3B")
    r3.add_break()
    p3.add_run("A4B")
    d.save(path)


def test_word_breaks_and_tabs_are_not_crossed(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _word(src / "a.docx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "A*B", "--to", "Z", "-o", str(out), "--allow-leftover")
    assert rc == 0, msg
    texts = [p.text for p in Document(str(out / "a.docx")).paragraphs]
    assert texts[0] == "A1\n2B"  # 改行をまたいで「A1<改行>2B」を 1 つにしない
    assert texts[2] == "Z\nZ"  # 改行の両側は別々に置換し、改行は残す
    body = _read(out / "a.docx", "word/document.xml")
    assert body.count("<w:br/>") == 2
    rules = tmp_path / "rules.csv"
    rules.write_text("置換前,置換後\n会員管理,X\n2B,2C\n", encoding="utf-8")
    out2 = tmp_path / "out2"
    rc, msg = run(str(src), "--rules", str(rules), "-o", str(out2), "--allow-leftover")
    assert rc == 0, msg
    texts = [p.text for p in Document(str(out2 / "a.docx")).paragraphs]
    assert texts[0] == "A1\n2C"
    assert texts[1] == "会員\t管理"  # タブをまたがない


def test_powerpoint_line_breaks_are_not_crossed(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tf = slide.shapes.title.text_frame
    tf.text = "A1\vB2"  # python-pptx は \v を a:br にする
    box = slide.shapes.add_textbox(0, 0, 100, 100).text_frame
    box.text = "A3B\vA4B"
    prs.save(src / "a.pptx")
    assert "<a:br" in _read(src / "a.pptx", "ppt/slides/slide1.xml")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "A*B", "--to", "Z", "-o", str(out), "--allow-leftover")
    assert rc == 0, msg
    shapes = Presentation(str(out / "a.pptx")).slides[0].shapes
    assert shapes.title.text_frame.text == "A1\vB2"  # 改行をまたがない
    assert shapes[1].text_frame.text == "Z\vZ"
    assert _read(out / "a.pptx", "ppt/slides/slide1.xml").count("<a:br") == 2
