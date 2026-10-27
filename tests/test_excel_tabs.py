"""Excel のセルの中のタブ・改行: ワイルドカード（* と ?）はまたがない（README「段落・セル内の改行とタブはまたぎません」）.

後退: Word・PowerPoint の改行・タブは別のランとして切っていたが、Excel のセルの文字列にはタブ（\\t）がそのまま入り、
`*`（`.*`）と `?`（`.`）がタブをまたいで一致していた（「A1<タブ>2B」が `A*B` で 1 つに置換され、タブが消えた）.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from test_office_replace import run

from office_replace.rules import Rule


def test_wildcards_do_not_match_tab_or_newline():
    assert Rule("A*B", "Z").apply("A1\t2B") == ("A1\t2B", 0)
    assert Rule("A?B", "Z").apply("A\tB") == ("A\tB", 0)
    assert Rule("A*B", "Z").apply("A1\r\n2B") == ("A1\r\n2B", 0)
    # タブの両側では別々に置換し、タブは残す
    assert Rule("A*B", "Z").apply("A1B\tA2B") == ("Z\tZ", 2)
    # 完全一致: タブを含むセルの全体には * で当たらない（Word の段落と同じ）
    assert Rule("A*B", "Z", whole=True).apply("A1\t2B") == ("A1\t2B", 0)
    # 正規表現（--regex）は書いたとおり（\t も . も利用者の指定に従う）
    assert Rule("A.*B", "Z", regex=True).apply("A1\t2B") == ("Z", 1)


def test_excel_cell_tab_is_not_crossed(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "A1\t2B"
    ws["A2"] = "A3B\tA4B"
    ws["A3"] = "A1\n2B"
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "A*B", "--to", "Z", "-o", str(out), "--allow-leftover")
    assert rc == 0, msg
    ws2 = load_workbook(out / "a.xlsx").active
    assert ws2["A1"].value == "A1\t2B"
    assert ws2["A2"].value == "Z\tZ"
    assert ws2["A3"].value == "A1\n2B"
