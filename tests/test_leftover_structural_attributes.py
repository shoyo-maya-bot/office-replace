"""残りの確認: XML の構造用の属性（色・大きさ・座標・ID など）は調べず、文字と意味のある属性（名前など）だけを調べる."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Emu
from test_leftover_sheet_names import _leftover_rows
from test_office_replace import report_of, run

from office_replace import leftover


def _files(src: Path) -> None:
    wb = Workbook()
    ws = wb.active
    for i in range(1, 60):
        ws.cell(i, 1, f"項目{i}")
        ws.cell(i, 2, i * 7)  # 50 に一致する数値は無い
        ws.cell(i, 1).fill = PatternFill("solid", fgColor="FF5050")
        ws.cell(i, 1).font = Font(size=50 if i == 3 else 11, color="505050")
    ws["C1"] = "単価 50 円"
    ws.column_dimensions["A"].width = 50
    ws.row_dimensions[50].height = 50
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=59))
    ws.add_chart(chart, "D50")
    wb.save(src / "a.xlsx")
    doc = Document()
    doc.add_paragraph("価格は 50 円")
    doc.save(src / "b.docx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "タイトル 50"
    slide.shapes.add_textbox(Emu(505050), Emu(5050500), Emu(1500000), Emu(500000)).text_frame.text = "x"
    prs.save(src / "c.pptx")


def test_structural_attributes_are_not_leftovers(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _files(src)
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "50", "--to", "60", "-o", str(out))
    assert rc == 0, msg  # 色（FF5050）・幅・座標・テーマなどに当たって 1 にならない
    assert not _leftover_rows(report_of(out).read_text(encoding="utf-8"))


def test_text_and_meaningful_attributes_are_still_checked():
    xml = (
        b'<r xmlns="urn:x"><sheet name="Shop50" sheetId="50"/><color rgb="FF5050"/>'
        b'<cNvPr id="50" name="Shape" descr="50 alt"/><off x="5050" y="0"/><t>50</t>'
        b'<alias val="50"/><sz val="50"/>'
        b'<Relationship Id="rId50" Target="file:///50/a.xlsx" TargetMode="External"/>'
        b'<Relationship Id="rId51" Target="sheet50.xml"/></r>'
    )
    places = [w for w, _ in leftover._strings(xml, None, "xl/workbook.xml")]
    assert "sheet@name" in places and "cNvPr@descr" in places and "alias@val" in places and "t" in places
    assert "Relationship@Target" in places  # 外部へのリンク先は調べる
    for structural in ("sheet@sheetId", "color@rgb", "cNvPr@id", "off@x", "sz@val", "Relationship@Id"):
        assert structural not in places
    targets = [v for w, v in leftover._strings(xml, None, "xl/workbook.xml") if w == "Relationship@Target"]
    assert targets == ["file:///50/a.xlsx"]  # 部品の中のリンク先（sheet50.xml）は調べない


def test_drawing_anchor_and_docprops_stats_are_not_checked():
    drawing = b'<w xmlns="urn:x"><from><col>49</col><colOff>50</colOff><row>50</row></from><t>a</t></w>'
    assert [(w, v) for w, v in leftover._strings(drawing, None, "xl/drawings/drawing1.xml")] == [("t", "a")]
    app = b'<Properties xmlns="urn:x"><TotalTime>50</TotalTime><Pages>50</Pages><Company>50</Company></Properties>'
    assert [(w, v) for w, v in leftover._strings(app, None, "docProps/app.xml")] == [("Company", "50")]
