"""テーブルの列名・ワイルドカードの最長一致・レポートの場所と伏せ字・ヘッダー/フッター・数式の文字列・置換表の読み込み."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from test_office_replace import _all_xml, _patch_zip, _read, report_of, run

from office_replace.processors import rename_table_columns, replace_literals
from office_replace.rules import LONG_MATCH, Rule, load

# --- ワイルドカード（Excel と同じ最長一致） ------------------------------------------------------------


def test_wildcard_star_is_greedy_like_excel():
    # 採点で再現した入力: 見積* → 御見積 が「御見積書A」になっていた（最短一致）
    assert Rule("見積*", "御見積").apply("見積書A") == ("御見積", 1)
    assert Rule("株式会社*支店", "（支店名）").apply("株式会社テスト 新宿支店・渋谷支店") == ("（支店名）", 1)
    assert Rule("担当: ?田", "X").apply("担当: 山田 担当: 田中") == ("X 担当: 田中", 1)


@pytest.mark.parametrize(
    ("before", "text", "expected"),
    [
        # 第 1 回の後退: 1 段落に 2 か所あると、1 か所目から 2 か所目の終わりまで 1 つにまとめて消えていた（間の「と」が消える）
        ("株式会社*支店", "株式会社A支店と株式会社B支店", ("XとX", 2)),
        ("株式会社*支店", "株式会社A支店と株式会社B支店です", ("XとXです", 2)),
        ("担当?*様", "担当A山田様、担当B田中様", ("X、X", 2)),
        # 頭（最初の * より前）が次に現れる手前で止め、2 か所目は別に当てる
        ("見積*", "見積書A、見積書B", ("XX", 2)),
        # 頭が無い（* から始まる）ときは止める目印が無いので、行の中では最長
        ("*支店", "A支店とB支店", ("X", 1)),
        ("abc*xyz", "ABC1xyz", ("ABC1xyz", 0)),  # -i なしは大文字小文字を区別
    ],
)
def test_wildcard_star_does_not_swallow_second_occurrence(before: str, text: str, expected: tuple[str, int]):
    assert Rule(before, "X").apply(text) == expected


def test_wildcard_star_ignore_case_stops_at_second_occurrence():
    assert Rule("abc*xyz", "X", ignore_case=True).apply("ABC1xyzとabc2XYZ") == ("XとX", 2)


def test_wildcard_star_stops_at_newline():
    # 段落・セル内の改行はまたがない（第 1 回は DOTALL で改行の先まで消していた）
    assert Rule("A*", "X").apply("A1\nA2") == ("X\nX", 2)
    assert Rule("株式会社*支店", "X").apply("株式会社A\n本社と支店") == ("株式会社A\n本社と支店", 0)
    # 置換前そのものに改行があるとき（置換表のセル内改行）は、その改行には一致する
    assert Rule("一行目\n二*", "X").apply("一行目\n二行目") == ("X", 1)


def test_wildcard_long_match_is_warned():
    r = Rule("株式会社*支店", "X")
    assert r.apply("株式会社" + "あ" * 10 + "支店") == ("X", 1)
    assert r.long_matches == []
    assert r.apply("株式会社" + "あ" * (LONG_MATCH + 1) + "支店") == ("X", 1)  # 置換はする. 警告用に長さを残す
    assert r.long_matches == [LONG_MATCH + 7]
    plain = Rule("株式会社" + "あ" * (LONG_MATCH + 1), "X")  # * の無いルールは長くても警告しない
    plain.apply(plain.before)
    assert plain.long_matches == []


def test_wildcard_two_occurrences_in_word_paragraph_and_excel_cell(tmp_path: Path):
    from docx import Document

    src = tmp_path / "in"
    src.mkdir()
    d = Document()
    p = d.add_paragraph("株式会社A支店と")
    p.add_run("株式会社B").bold = True  # 書式の切れ目をまたいでも同じ
    p.add_run("支店にご案内")
    d.add_paragraph("次の段落 株式会社C支店")
    d.save(src / "a.docx")
    wb = Workbook()
    wb.active["A1"] = "株式会社A支店と株式会社B支店"
    wb.active["A2"] = "株式会社C\n支店"  # セル内の改行はまたがない
    wb.save(src / "b.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "株式会社*支店", "--to", "（社名）", "-o", str(out), "--allow-leftover")
    assert rc == 0, msg
    texts = [p.text for p in Document(str(out / "a.docx")).paragraphs]
    assert texts == ["（社名）と（社名）にご案内", "次の段落 （社名）"]
    ws = load_workbook(out / "b.xlsx").active
    assert ws["A1"].value == "（社名）と（社名）"
    assert ws["A2"].value == "株式会社C\n支店"


def test_wildcard_long_match_warning_in_report(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active["A1"] = "株式会社" + "あ" * (LONG_MATCH + 5) + "支店"
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "株式会社*支店", "--to", "X", "-o", str(out), "--mask-report")
    assert rc == 0, msg  # 警告は終了コードに影響しない
    assert "警告: 1 件" in msg
    md = report_of(out).read_text(encoding="utf-8")
    assert f"ルール 1 の * が {LONG_MATCH} 文字より長く一致" in md
    assert "株式会社" not in md  # 警告に置換前の文字列は入れない（--mask-report でも漏れない）


# --- 置換表の読み込み ---------------------------------------------------------------------------------


def test_csv_cell_with_newline(tmp_path: Path):
    p = tmp_path / "r.csv"
    p.write_text('置換前,置換後\n"一行目\n二行目",X\n次,Y\n', encoding="utf-8")
    rules = load(p)
    assert [(r.before, r.after) for r in rules] == [("一行目\n二行目", "X"), ("次", "Y")]


def test_xlsx_rules_are_read_as_displayed(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["置換前", "置換後"])
    ws.append([1, "コード001"])
    ws["A2"].number_format = "000"  # 表示は 001
    ws.append([2.0, 3])  # 2.0 は「2」
    wb.save(tmp_path / "r.xlsx")
    rules = load(tmp_path / "r.xlsx")
    assert [(r.before, r.after) for r in rules] == [("001", "コード001"), ("2", "3")]


# --- テーブル（ListObject）の列名と構造化参照 -----------------------------------------------------------


def _table_book(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["架空商会コード", "金額"])
    ws.append(["A", 100])
    ws.append(["B", 200])
    ws["D1"] = "=SUM(テーブル1[金額])"
    ws["D2"] = "=COUNTA(テーブル1[架空商会コード])"
    ws["D3"] = "=テーブル1[[#Headers],[架空商会コード]]"
    tab = Table(displayName="テーブル1", ref="A1:B3")
    ws.add_table(tab)
    wb.save(path)

    def add_name(data: bytes) -> bytes:
        return data.replace(
            b"</sheets>",
            '</sheets><definedNames><definedName name="コード列">テーブル1[架空商会コード]</definedName></definedNames>'.encode(),
        )

    _patch_zip(path, {"xl/workbook.xml": add_name})


def test_table_column_names_follow_header_cells(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _table_book(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "取引先", "-o", str(out))
    assert rc == 0, msg
    res = out / "a.xlsx"
    with zipfile.ZipFile(res) as z:
        table_part = next(n for n in z.namelist() if n.startswith("xl/tables/"))
    table = _read(res, table_part)
    assert 'name="取引先コード"' in table and "架空" not in table
    sheet = _read(res, "xl/worksheets/sheet1.xml")
    assert "COUNTA(テーブル1[取引先コード])" in sheet
    assert "テーブル1[[#Headers],[取引先コード]]" in sheet
    assert "SUM(テーブル1[金額])" in sheet
    assert '<definedName name="コード列">テーブル1[取引先コード]</definedName>' in _read(res, "xl/workbook.xml")
    assert "架空" not in _all_xml(res)
    wb = load_workbook(res)
    t = wb["一覧"].tables["テーブル1"]
    assert [c.name for c in t.tableColumns] == ["取引先コード", "金額"]
    assert wb["一覧"]["A1"].value == "取引先コード"  # 見出しセルと一致
    md = report_of(out).read_text(encoding="utf-8")
    assert "テーブル「テーブル1」の列名" in md and "構造化参照" in md


def test_table_column_duplicate_is_error(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _table_book(src / "a.xlsx")
    rc, _ = run(str(src), "--from", "架空商会コード", "--to", "金額", "-o", str(tmp_path / "out"))
    assert rc == 1
    assert "列名が置換で重複" in report_of(tmp_path / "out").read_text(encoding="utf-8")


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("SUM(T1[旧])", "SUM(T1[新'[1']])"),
        ("T1[[#This Row],[旧]]*2", "T1[[#This Row],[新'[1']]]*2"),
        ("SUM(T1[[旧]:[他]])", "SUM(T1[[新'[1']]:[他]])"),
        ('T2[旧]&"T1[旧]"', 'T2[旧]&"T1[旧]"'),  # 別のテーブル・文字列の中は変えない
        ("t1[旧]", "t1[新'[1']]"),  # テーブル名・列名は大文字小文字を区別しない
        ("T1['#旧]", "T1['#旧]"),  # 列名「#旧」は別の列
    ],
)
def test_rename_table_columns(formula: str, expected: str):
    assert rename_table_columns(formula, {"t1": {"旧": "新[1]"}}) == expected


# --- レポートの場所・伏せ字・終了コード ---------------------------------------------------------------


def test_report_is_outside_output_and_can_be_masked(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active.title = "架空商会向け"
    wb.active["A1"] = "架空商会 御中"
    wb.save(src / "架空商会_見積.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "--target", "all", "-o", str(out), "--mask-report")
    assert rc == 0, msg
    # 出力フォルダには置換後のファイルだけ（レポートは出力フォルダの外）
    assert sorted(p.name for p in out.iterdir()) == ["〇〇_見積.xlsx"]
    md = report_of(out).read_text(encoding="utf-8")
    assert "架空" not in md and "（伏せ字）" in md
    rep = load_workbook(report_of(out).with_suffix(".xlsx"))
    assert all("架空" not in str(c.value) for ws in rep.worksheets for row in ws.iter_rows() for c in row)
    # 伏せないときは、置換前が入る（だから出力フォルダの外に置く）
    run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(tmp_path / "out2"))
    assert "架空商会" in report_of(tmp_path / "out2").read_text(encoding="utf-8")


def test_report_inside_output_warns(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    Workbook().save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空", "--to", "〇〇", "-o", str(out), "--report", str(out / "r.md"))
    assert rc == 0 and "警告" in msg and (out / "r.md").exists()


# --- ヘッダー・フッターと数式の中の文字列 -------------------------------------------------------------


def test_excel_header_footer_is_replaced(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "本文"
    ws.oddHeader.center.text = "架空商会 御中"
    ws.oddFooter.right.text = "架空商会 社外秘"
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(out))
    assert rc == 0, msg
    assert "架空" not in _all_xml(out / "a.xlsx")
    ws2 = load_workbook(out / "a.xlsx").active
    assert ws2.oddHeader.center.text == "〇〇 御中" and ws2.oddFooter.right.text == "〇〇 社外秘"
    assert "ヘッダー" in report_of(out).read_text(encoding="utf-8")


def test_replace_literals():
    rules = [Rule("架空", "〇〇")]
    assert replace_literals('="架空商会"&A1&"様"', rules) == ('="〇〇商会"&A1&"様"', 1)
    assert replace_literals('=架空!A1&"x""架空"""', rules) == ('=架空!A1&"x""〇〇"""', 1)  # 引用符の中だけ


def test_formula_target_replaces_string_literals(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "見出し"
    ws["A2"] = '="架空商会"&"様"'
    wb.save(src / "a.xlsx")
    _patch_zip(
        src / "a.xlsx",
        {
            # 計算結果（キャッシュ）として置換前の文字列を持たせる
            "xl/worksheets/sheet1.xml": lambda d: d.replace(b"<v></v>", "<v>架空商会様</v>".encode()).replace(
                b'<c r="A2">', b'<c r="A2" t="str">'
            )
        },
    )
    # 既定（文字列だけ）では、数式の中は残り（理由つき）
    rc, _ = run(str(src), "--from", "架空商会", "--to", "〇〇", "-o", str(tmp_path / "out1"))
    assert rc == 1
    assert "--target formula" in report_of(tmp_path / "out1").read_text(encoding="utf-8")
    # --target text,formula で数式の中の文字列も置換し、計算結果は開いたときに計算し直させる
    out = tmp_path / "out2"
    rc, msg = run(str(src), "--from", "架空商会", "--to", "〇〇", "--target", "text,formula", "-o", str(out))
    assert rc == 0, msg
    assert "架空" not in _all_xml(out / "a.xlsx")
    assert load_workbook(out / "a.xlsx").active["A2"].value == '="〇〇"&"様"'
    assert 'fullCalcOnLoad="1"' in _read(out / "a.xlsx", "xl/workbook.xml")
    assert "<v>" not in _read(out / "a.xlsx", "xl/worksheets/sheet1.xml")  # 古い計算結果は消す
    assert "数式の文字列" in report_of(out).read_text(encoding="utf-8")
