"""テーブルの列名が置換で重複するとき、壊れたファイル（Excel が修復を求める）を出力しない."""

from __future__ import annotations

from pathlib import Path

from test_office_replace import report_of, run
from test_tables_formulas_report import _table_book


def test_duplicate_table_column_is_not_written(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _table_book(src / "a.xlsx")
    (src / "b.xlsx").write_bytes((src / "a.xlsx").read_bytes())
    out = tmp_path / "out"
    rc, _ = run(str(src), "--from", "架空商会コード", "--to", "金額", "-o", str(out))
    assert rc == 1
    assert not (out / "a.xlsx").exists() and not (out / "b.xlsx").exists()
    md = report_of(out).read_text(encoding="utf-8")
    assert "列名が置換で重複" in md and "出力しません" in md


def test_duplicate_table_column_in_dry_run_is_error(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    _table_book(src / "a.xlsx")
    out = tmp_path / "out"
    rc, _ = run(str(src), "--from", "架空商会コード", "--to", "金額", "-o", str(out), "--dry-run")
    assert rc == 1
    assert not (out / "a.xlsx").exists()


def test_other_files_are_still_written(tmp_path: Path):
    from openpyxl import Workbook, load_workbook

    src = tmp_path / "in"
    src.mkdir()
    _table_book(src / "a.xlsx")
    wb = Workbook()
    wb.active["A1"] = "架空商会コード一覧"
    wb.save(src / "c.xlsx")
    out = tmp_path / "out"
    rc, _ = run(str(src), "--from", "架空商会コード", "--to", "金額", "-o", str(out))
    assert rc == 1
    assert not (out / "a.xlsx").exists()
    assert load_workbook(out / "c.xlsx").active["A1"].value == "金額一覧"
