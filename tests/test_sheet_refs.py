"""シート名の変更に合わせた数式の参照: 3D 参照の両端を直し、外部のブックの参照は変えない."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from test_office_replace import _read, run

from office_replace.processors import rename_sheet_refs

R = {"別": "新別", "集計": "合計", "It's": "Y"}


@pytest.mark.parametrize(
    ("formula", "expected", "count"),
    [
        ("SUM(集計:別!A1)", "SUM('合計:新別'!A1)", 1),  # 3D 参照（引用符なし）の両端
        ("SUM(先頭:別!A1)", "SUM('先頭:新別'!A1)", 1),  # 3D 参照の終わりだけ
        ("SUM('集計:別'!A1)", "SUM('合計:新別'!A1)", 1),  # 3D 参照（引用符あり）
        ("[1]集計!A1", "[1]集計!A1", 0),  # 外部のブックは変えない
        ("'[1]集計'!A1+[2]別!B2", "'[1]集計'!A1+[2]別!B2", 0),
        ("集計!A1+別!B2", "'合計'!A1+'新別'!B2", 2),
        ("'It''s'!A1", "'Y'!A1", 1),
        ('INDIRECT("集計!A1")', 'INDIRECT("集計!A1")', 0),  # 文字列の中は変えない
        ("旧集計!A1", "旧集計!A1", 0),  # 名前の一部では一致しない
    ],
)
def test_rename_sheet_refs(formula: str, expected: str, count: int):
    assert rename_sheet_refs(formula, R) == (expected, count)


def test_sheet_rename_fixes_3d_reference_but_not_external(tmp_path: Path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "集計"
    wb.create_sheet("別")["A1"] = 1
    ws["B1"] = "=SUM(集計:別!A1)"
    ws["B2"] = "=[1]集計!A1"
    wb.save(src / "a.xlsx")
    out = tmp_path / "out"
    rc, msg = run(str(src), "--from", "別", "--to", "新別", "--target", "sheet", "-o", str(out))
    assert rc == 0, msg
    s1 = _read(out / "a.xlsx", "xl/worksheets/sheet1.xml")
    assert "<f>SUM('集計:新別'!A1)</f>" in s1
    out2 = tmp_path / "out2"
    rc, msg = run(str(src), "--from", "集計", "--to", "合計", "--target", "sheet", "-o", str(out2))
    assert rc == 0, msg
    s1 = _read(out2 / "a.xlsx", "xl/worksheets/sheet1.xml")
    assert "<f>SUM('合計:別'!A1)</f>" in s1
    assert "<f>[1]集計!A1</f>" in s1  # 外部のブックのシートは書き換えない
