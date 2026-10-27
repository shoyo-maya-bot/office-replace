"""--rules と一緒に付けた --to / --target / --whole は使われない. 黙って無視せず警告する（互換のため終了コードは変えない）."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from test_office_replace import run


def _input(tmp_path: Path) -> Path:
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active["A1"] = "旧社名"
    wb.save(src / "a.xlsx")
    return src


def test_rules_with_from_only_options_warns(tmp_path: Path):
    src = _input(tmp_path)
    rules = tmp_path / "rules.csv"
    rules.write_text("置換前,置換後\n旧社名,新社名\n", encoding="utf-8")
    rc, msg = run(str(src), "--rules", str(rules), "--to", "別", "--target", "all", "--whole", "-o", str(tmp_path / "out"))
    assert rc == 0, msg
    assert "警告: --to / --target / --whole は --from と組で使うオプションです" in msg
    # 置換は置換表のとおり（--to は使わない）
    assert load_workbook(tmp_path / "out" / "a.xlsx").active["A1"].value == "新社名"


def test_no_warning_without_from_only_options(tmp_path: Path):
    src = _input(tmp_path)
    rules = tmp_path / "rules.csv"
    rules.write_text("置換前,置換後\n旧社名,新社名\n", encoding="utf-8")
    rc, msg = run(str(src), "--rules", str(rules), "-o", str(tmp_path / "out"))
    assert rc == 0, msg
    assert "--from と組で使う" not in msg
    rc, msg = run(str(src), "--from", "旧社名", "--target", "text", "-o", str(tmp_path / "out2"))
    assert rc == 0, msg
    assert "--from と組で使う" not in msg
    # --to を省くと置換後は空（README のとおり）
    assert load_workbook(tmp_path / "out2" / "a.xlsx").active["A1"].value in (None, "")
